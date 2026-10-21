"""QGIS 4 runtime checks for the modeless Designer-based Survey workflow."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from qgis.PyQt.QtWidgets import QMainWindow, QDockWidget, QStackedWidget, QMenu, QMessageBox
from qgis.PyQt.uic import loadUi
from qgis.PyQt.QtCore import Qt, QDate, QSettings
from qgis.core import (QgsApplication, QgsVectorLayer, QgsFeature, QgsGeometry,
                       QgsCoordinateReferenceSystem, QgsCoordinateTransform,
                       QgsCoordinateTransformContext, QgsPointXY, QgsProject,
                       Qgis)

from geoflow_connector.ui.survey_panel import SurveyPanel, SurveyValidationError
from geoflow_connector.ui.survey_import_dialog import ImportDialog
from geoflow_connector.app.survey_integration import SurveyIntegrationMixin
from geoflow_connector.app.integration import IntegrationEvents
from geoflow_connector.api.references import normalize_catalog
from geoflow_connector.survey.code_mapping import (
    depth_create_key, facility_identity, read_depth_create_targets,
    read_mappings, read_storage_codes, storage_code_key,
    write_depth_create_targets, write_mappings, write_storage_codes,
)
from geoflow_connector.survey.code_normalizer import (
    normalize_raw_code, normalization_key, write_code_aliases,
    read_code_rules, write_code_rules, write_code_rule_rows,
)


APP = QgsApplication.instance() or QgsApplication([], False)
APP.initQgis()
PROJECT = "10000000-0000-4000-8000-000000000001"
SOURCE = "30000000-0000-4000-8000-000000000001"
SURVEY = "50000000-0000-4000-8000-000000000001"
TARGET = "60000000-0000-4000-8000-000000000001"
MAPPING = "40000000-0000-4000-8000-000000000001"


def fixture(kind="Point", wkt="POINT (127 36)"):
    fields = "field=id:string&field=FTR_CDE:string"
    layer = QgsVectorLayer(f"{kind}?crs=EPSG:4326&{fields}", "시설물", "memory")
    layer.setCustomProperty("geoflow/standard_name", "TEST_LAYER")
    feature = QgsFeature(layer.fields()); feature["id"] = TARGET
    feature.setGeometry(QgsGeometry.fromWkt(wkt)); layer.dataProvider().addFeature(feature)
    layer.selectByIds([next(layer.getFeatures()).id()])
    window = QMainWindow(); dock = QDockWidget(window); stack = QStackedWidget(); dock.setWidget(stack)
    client = Mock(base_url="https://example.invalid", session_user={"user_id": PROJECT})
    source = {"id": SOURCE, "original_file_name": "survey.csv", "version": 1, "is_active": True}
    point = {"id": SURVEY, "source_id": SOURCE, "source_row_id": "P001", "raw_code": "M",
             "longitude": 127.0, "latitude": 36.0, "raw_z": 12.3}
    def get(url):
        if "survey-sources/" in url: return {"results": [source]}
        if "survey-points/" in url: return {"results": [point]}
        return {"results": []}
    client.get_json.side_effect = get
    client.post_json.return_value = {"ok": True}
    work = SimpleNamespace(_dirty_page=lambda: None, pages={}, on_feature_selected=Mock())
    iface = SimpleNamespace(mainWindow=lambda: window, activeLayer=lambda: layer,
                            mapCanvas=lambda: SimpleNamespace(refresh=Mock()))
    plugin = SimpleNamespace(
        active_client=client, active_context={"project_id": PROJECT, "client_id": PROJECT,
            "changeset_supported": True, "delta_url": "/delta", "package_path": "unused",
            "manifest": {}},
        integration_state=lambda: {"ready": True, "can_write": True, "project_id": PROJECT, "epoch": 1},
        _managed_layers=lambda: [layer], _sync_in_progress=False, _sync_active_project=Mock(),
        _pull_clean_snapshot_delta=Mock(return_value=0),
        _ensure_work=lambda: None, work=work, iface=iface, dock=dock,
        _reference_service=SimpleNamespace(catalog={"groups": [{"code_group_key": "FTR", "values": [
            {"code": "SA100", "label": "맨홀", "sort_order": 0}]}], "bindings": [
            {"standard_name": "TEST_LAYER", "field_name": "FTR_CDE", "code_group_key": "FTR"}]})
    )
    return window, layer, plugin


class PopupTests(unittest.TestCase):
    def setUp(self):
        self.window, self.layer, self.plugin = fixture()
        self.settings_temp = tempfile.TemporaryDirectory()
        settings_file = str(Path(self.settings_temp.name) / "popup.ini")
        self.patches = [patch(
            "geoflow_connector.ui.survey_panel.QSettings",
            side_effect=lambda: QSettings(
                settings_file, QSettings.Format.IniFormat
            ),
        )]
        for item in self.patches: item.start()

    def tearDown(self):
        for item in self.patches: item.stop()
        self.window.close(); self.window.deleteLater()
        self.settings_temp.cleanup()

    def select_source(self, popup, source_id=SOURCE):
        index = popup.sourceCombo.findData(source_id)
        self.assertGreater(index, 0)
        popup.sourceCombo.setCurrentIndex(index)
        APP.processEvents()
        return index

    def test_modeless_ui_points_mapping_and_debug(self):
        popup = SurveyPanel(self.plugin, self.window); popup.show(); APP.processEvents()
        self.assertFalse(popup.isModal()); self.assertEqual(popup.sourceCombo.count(), 2)
        self.assertEqual(popup.sourceCombo.currentIndex(), 0)
        self.assertEqual(popup.sourceCombo.currentText(), "측량원시 선택")
        self.assertEqual(popup.featureTable.rowCount(), 0)
        self.assertFalse(popup.createBatchButton.isEnabled())
        self.select_source(popup)
        self.assertFalse(popup.debugText.isVisible())
        popup.points(); self.assertEqual(popup.pointsTable.rowCount(), 1)
        popup.render_mapping(); self.assertEqual(popup.mappingTable.rowCount(), 1)
        self.assertEqual(
            [popup.mappingTabs.tabText(i) for i in range(popup.mappingTabs.count())],
            ["시설물 생성", "코드 정규화", "코드/시설물 매핑", "저장코드"],
        )
        self.assertEqual(popup.mappingTabs.indexOf(popup.depthCreateTab), -1)
        self.assertEqual(popup.pages.currentWidget(), popup.mappingPage)
        self.assertEqual(popup.mappingTabs.currentWidget(), popup.featurePage)
        for name in ("pointsButton", "mappingButton", "featureButton",
                     "linksButton", "reapplyButton"):
            self.assertTrue(
                not hasattr(popup, name) or getattr(popup, name).isHidden()
            )
        self.assertTrue(popup.debugGroup.isHidden())
        combo = popup.mappingTable.cellWidget(0, 1); combo.setCurrentIndex(1)
        popup.mappingTable.setCurrentCell(0, 0); popup.apply_same_code()
        self.assertEqual(popup.pointsTable.item(0, 3).text(), "맨홀")
        popup.pages.setCurrentWidget(popup.mappingPage)
        popup.mappingTabs.setCurrentWidget(popup.normalizationTab)
        popup.resize(900, 700); APP.processEvents()
        artifact = (Path(__file__).resolve().parents[1] / "test_artifacts" /
                    "survey" / "code-normalization-4tabs-qgis422.png")
        artifact.parent.mkdir(parents=True, exist_ok=True)
        self.assertTrue(popup.grab().save(str(artifact)))
        self.assertFalse(popup.grab().isNull())
        popup.close_service(); popup.close(); popup.deleteLater()

    def test_no_sources_shows_disabled_placeholder(self):
        self.plugin.active_client.get_json.side_effect = lambda _url: {"results": []}
        popup = SurveyPanel(self.plugin, self.window)
        self.assertEqual(popup.sourceCombo.count(), 1)
        self.assertEqual(popup.sourceCombo.currentText(), "측량원시 선택")
        self.assertFalse(popup.sourceCombo.isEnabled())
        self.assertEqual(
            popup.sourceCombo.toolTip(), "등록된 측량원시가 없습니다."
        )
        self.assertEqual(popup.featureTable.rowCount(), 0)
        self.assertFalse(popup.createBatchButton.isEnabled())
        popup.close(); popup.deleteLater()

    def test_source_selection_switch_refresh_and_restart_load_all_points(self):
        source_b = "30000000-0000-4000-8000-000000000002"
        source_c = "30000000-0000-4000-8000-000000000003"
        sources = [
            {"id": SOURCE, "original_file_name": "A.csv", "version": 1},
            {"id": source_b, "original_file_name": "B.csv", "version": 1},
            {"id": source_c, "original_file_name": "C.csv", "version": 1},
        ]

        def point_rows(source_id, count):
            return [{
                "id": f"50000000-0000-4000-8001-{number:012d}",
                "source_id": source_id, "source_row_id": f"P{number:03d}",
                "survey_no": f"261017-{number:03d}", "code": "M",
                "longitude": 127.0, "latitude": 36.0,
                "x": 235000.0 + number, "y": 413000.0 + number,
                "z": 50.0,
            } for number in range(1, count + 1)]

        rows_by_source = {
            SOURCE: point_rows(SOURCE, 22),
            source_b: point_rows(source_b, 73),
            source_c: [],
        }

        def get(url):
            if "survey-sources/" in url:
                return {"results": list(sources)}
            if "survey-points/" in url:
                source_id = next((key for key in rows_by_source if key in url), None)
                return {"results": list(rows_by_source.get(source_id, []))}
            return {"results": []}

        self.plugin.active_client.get_json.side_effect = get
        popup = SurveyPanel(self.plugin, self.window)
        self.assertEqual(popup.sourceCombo.count(), 4)
        self.assertEqual(popup.sourceCombo.currentIndex(), 0)
        self.assertEqual(popup.featureTable.rowCount(), 0)

        popup.sourceCombo.setCurrentIndex(popup.sourceCombo.findData(SOURCE))
        self.assertEqual(popup.featureTable.rowCount(), 22)
        self.assertTrue(all(
            popup.featureTable.item(row, 1).data(Qt.ItemDataRole.UserRole)["source_id"]
            == SOURCE for row in range(22)
        ))

        popup.sourceCombo.setCurrentIndex(popup.sourceCombo.findData(source_b))
        self.assertEqual(popup.featureTable.rowCount(), 73)
        self.assertTrue(all(
            popup.featureTable.item(row, 1).data(Qt.ItemDataRole.UserRole)["source_id"]
            == source_b for row in range(73)
        ))

        rows_by_source[source_b] = point_rows(source_b, 74)
        popup.sourceRefreshButton.click(); APP.processEvents()
        self.assertEqual(popup.sourceCombo.currentData(), source_b)
        self.assertEqual(popup.featureTable.rowCount(), 74)

        sources[:] = [row for row in sources if row["id"] != source_b]
        popup.sourceRefreshButton.click(); APP.processEvents()
        self.assertEqual(popup.sourceCombo.currentIndex(), 0)
        self.assertIsNone(popup.sourceCombo.currentData())
        self.assertEqual(popup.featureTable.rowCount(), 0)
        self.assertFalse(popup.createBatchButton.isEnabled())

        restarted = SurveyPanel(self.plugin, self.window)
        self.assertEqual(restarted.sourceCombo.currentIndex(), 0)
        restarted.sourceCombo.setCurrentIndex(
            restarted.sourceCombo.findData(SOURCE)
        )
        self.assertEqual(restarted.featureTable.rowCount(), 22)
        restarted.close(); restarted.deleteLater()
        popup.close(); popup.deleteLater()

    def test_project_settings_autoload_without_source_and_survive_switch(self):
        seed = SurveyPanel(self.plugin, self.window)
        normalization_settings, normalization_key_value = seed._normalization_settings()
        mapping_settings, mapping_key_value = seed._standard_mapping_settings()
        storage_settings, storage_key_value = seed._storage_settings()
        depth_settings, depth_key_value = seed._depth_settings()
        identity = facility_identity("TEST_LAYER", "맨홀")
        write_code_rule_rows(normalization_settings, normalization_key_value, [{
            "business_type": "", "raw_code": "mh",
            "survey_code": "M", "enabled": True,
        }])
        write_mappings(mapping_settings, mapping_key_value, {"M": {
            "facility_id": identity, "facility_label": "맨홀",
            "layer": "TEST_LAYER",
        }})
        write_storage_codes(storage_settings, storage_key_value, {identity: {
            "facility_id": identity, "facility_label": "맨홀",
            "standard_name": "TEST_LAYER", "storage_field": "ftr_cde",
            "storage_code": "SA100", "code_label": "맨홀",
        }})
        write_depth_create_targets(
            depth_settings, depth_key_value, {("TEST_LAYER", "SA100")}
        )
        seed.close(); seed.deleteLater()

        with patch(
            "geoflow_connector.ui.survey_panel.QgsMessageLog.logMessage"
        ) as log:
            popup = SurveyPanel(self.plugin, self.window)
        self.assertIsNone(popup.sourceCombo.currentData())
        self.assertEqual(popup.normalizationTable.rowCount(), 2)
        self.assertIn(
            "MH", [popup.normalizationTable.item(row, 1).text()
                   for row in range(popup.normalizationTable.rowCount())]
        )
        self.assertEqual(popup.mappingTable.rowCount(), 1)
        self.assertEqual(popup.mappingTable.item(0, 0).text(), "M")
        storage_rows = [
            row for row in range(popup.storageCodeTable.rowCount())
            if popup.storageCodeTable.item(row, 0).data(
                Qt.ItemDataRole.UserRole
            ) == identity
        ]
        self.assertEqual(len(storage_rows), 1)
        storage_row = storage_rows[0]
        self.assertEqual(
            popup.storageCodeTable.cellWidget(storage_row, 2).currentData()[0],
            "SA100",
        )
        self.assertTrue(
            popup.storageCodeTable.cellWidget(storage_row, 3).isChecked()
        )
        messages = [call.args[0] for call in log.call_args_list]
        self.assertTrue(any(
            "[SURVEY_SETTINGS_LOAD] normalization_count=1" in message
            for message in messages
        ))
        self.assertTrue(any(
            "[SURVEY_SETTINGS_LOAD] mapping_count=1" in message
            for message in messages
        ))

        before = (
            popup.normalizationTable.rowCount(), popup.mappingTable.rowCount(),
            popup.storageCodeTable.rowCount(),
        )
        self.select_source(popup)
        self.assertEqual(before, (
            popup.normalizationTable.rowCount(), popup.mappingTable.rowCount(),
            popup.storageCodeTable.rowCount(),
        ))
        popup.sourceCombo.setCurrentIndex(0)
        self.assertEqual(before, (
            popup.normalizationTable.rowCount(), popup.mappingTable.rowCount(),
            popup.storageCodeTable.rowCount(),
        ))

        restarted = SurveyPanel(self.plugin, self.window)
        self.assertIsNone(restarted.sourceCombo.currentData())
        self.assertEqual(restarted.normalizationTable.rowCount(), 2)
        self.assertEqual(restarted.mappingTable.rowCount(), 1)
        restarted.close(); restarted.deleteLater()
        popup.close(); popup.deleteLater()

    def test_apply_same_code_without_current_row_is_ui_validation(self):
        popup = SurveyPanel(self.plugin, self.window); self.select_source(popup); popup.render_mapping()
        popup.mappingTable.selectionModel().clear()
        popup.apply_same_code()
        self.assertEqual(
            popup.statusLabel.text(),
            "먼저 적용 기준이 될 측량점 행을 선택하세요.",
        )
        self.assertTrue(popup.applyCodeButton.isEnabled())
        popup.close(); popup.deleteLater()

    def test_normalization_uses_catalog_l2_and_blank_row_editor(self):
        self.plugin._reference_service.catalog["business_types"] = [
            {"code": "WATER", "label": "상수"}
        ]
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup); popup.render_mapping()
        self.assertEqual(
            [popup.normalizationTable.horizontalHeaderItem(i).text()
             for i in range(popup.normalizationTable.columnCount())],
            ["업무유형", "원시코드", "Survey Code", "사용"],
        )
        before = popup.normalizationTable.rowCount()
        popup.add_normalization_rule()
        self.assertEqual(popup.normalizationTable.rowCount(), before + 1)
        row = before
        business = popup.normalizationTable.cellWidget(row, 0)
        self.assertGreaterEqual(business.findData("WATER"), 0)
        self.assertEqual(popup.normalizationTable.item(row, 1).text(), "")
        self.assertEqual(popup.normalizationTable.item(row, 2).text(), "")
        popup.close(); popup.deleteLater()

    def test_reference_catalog_preserves_only_active_project_l2(self):
        catalog = normalize_catalog({
            "ok": True, "groups": [], "bindings": [],
            "business_types": [
                {"code": "WATER", "label": "상수", "sort_order": 2},
                {"code": "SEWER", "label": "하수", "sort_order": 1},
                {"code": "ROAD", "label": "도로", "active": False},
            ],
        })
        self.assertEqual(
            [(row["code"], row["label"]) for row in catalog["business_types"]],
            [("SEWER", "하수"), ("WATER", "상수")],
        )
        self.assertEqual(catalog["business_type_source"], "catalog.business_types")

    def test_l2_specific_same_raw_code_survives_reopen(self):
        self.plugin._reference_service.catalog["business_types"] = [
            {"code": "WATER", "label": "상수"},
            {"code": "SEWER", "label": "하수"},
        ]
        seed = SurveyPanel(self.plugin, self.window)
        settings, key = seed._normalization_settings()
        write_code_rule_rows(settings, key, [
            {"business_type": "WATER", "raw_code": "J",
             "survey_code": "J", "enabled": True},
            {"business_type": "SEWER", "raw_code": "J",
             "survey_code": "MH", "enabled": True},
        ])
        seed.close(); seed.deleteLater()
        popup = SurveyPanel(self.plugin, self.window)
        actual = set()
        for row in range(popup.normalizationTable.rowCount()):
            if popup.normalizationTable.item(row, 1).text() == "J":
                actual.add((
                    popup.normalizationTable.cellWidget(row, 0).currentData(),
                    popup.normalizationTable.item(row, 2).text(),
                ))
        self.assertEqual(actual, {("WATER", "J"), ("SEWER", "MH")})
        popup.close(); popup.deleteLater()

    def test_import_refreshes_project_wide_new_codes_without_reopen(self):
        popup = SurveyPanel(self.plugin, self.window)
        new_source = "30000000-0000-4000-8000-000000000099"
        codes = ("DCIP100", "J150", "J200", "P200-150")
        new_points = [{
            "id": f"50000000-0000-4000-8000-{index:012d}",
            "source_id": new_source, "source_row_id": f"N{index}",
            "raw_code": code, "code": code, "survey_no": f"261017-{index:03d}",
            "type": "GNSS", "longitude": 127.0, "latitude": 36.0,
            "x": 1.0, "y": 2.0, "z": 3.0,
        } for index, code in enumerate(codes, 1)]
        old_source = {"id": SOURCE, "original_file_name": "old.csv",
                      "version": 1, "is_active": True}
        created_source = {"id": new_source, "original_file_name": "new.csv",
                          "version": 1, "is_active": True}
        old_point = {"id": SURVEY, "source_id": SOURCE,
                     "source_row_id": "P001", "raw_code": "M", "code": "M"}

        def get(url):
            if "survey-sources/" in url:
                return {"results": [old_source, created_source]}
            if "survey-points/" in url:
                if new_source in url:
                    return {"results": new_points}
                if SOURCE in url:
                    return {"results": [old_point]}
                return {"results": [old_point, *new_points]}
            return {"results": []}

        self.plugin.active_client.get_json.side_effect = get
        popup.controller.import_source_file = Mock(return_value={
            "source_id": new_source, "created": 4, "updated": 0,
            "local_count": 4, "response_points": new_points,
        })
        popup.controller.refresh_imported_points = Mock(return_value=True)
        self.plugin.active_client.reset_mock()
        with patch(
            "geoflow_connector.ui.survey_panel._import_readback_verification",
            return_value={"ok": True, "checked_count": 4,
                          "issues": [], "sample": []},
        ):
            popup._submit_import(
                {"survey_type": "GNSS", "points": [
                    {"source_row_id": row["source_row_id"],
                     "raw_code": row["raw_code"]} for row in new_points
                ]}, Path("new.csv"), Mock(),
            )
        self.assertEqual(popup.sourceCombo.currentData(), new_source)
        self.assertEqual(popup.featureTable.rowCount(), 4)
        normalization_codes = {
            popup.normalizationTable.item(row, 1).text()
            for row in range(popup.normalizationTable.rowCount())
        }
        mapping_codes = {
            popup.mappingTable.item(row, 0).text()
            for row in range(popup.mappingTable.rowCount())
        }
        self.assertTrue(set(codes).issubset(normalization_codes))
        self.assertTrue(set(codes).issubset(mapping_codes))
        self.assertIn("M", mapping_codes)
        point_urls = [call.args[0] for call in
                      self.plugin.active_client.get_json.call_args_list
                      if "survey-points/" in call.args[0]]
        self.assertTrue(any("source_id=" not in url for url in point_urls))
        popup.close(); popup.deleteLater()

    def test_mapping_and_storage_show_manifest_layer_label(self):
        self.plugin.active_context["manifest"] = {"layers": [{
            "standard_name": "TEST_LAYER", "label": "테스트 시설물",
        }]}
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup); popup.render_mapping()
        combo = popup.mappingTable.cellWidget(0, 1)
        self.assertIn("테스트 시설물", combo.itemText(1))
        self.assertEqual(popup.storageCodeTable.item(0, 1).text(), "테스트 시설물")
        self.assertEqual(
            popup.storageCodeTable.item(0, 1).data(Qt.ItemDataRole.UserRole),
            "TEST_LAYER",
        )
        popup.close(); popup.deleteLater()

    def test_run_logs_validation_message_traceback_and_context(self):
        popup = SurveyPanel(self.plugin, self.window)
        def start_selected():
            raise SurveyValidationError(
                "storage_code_missing", "저장 코드가 설정되지 않았습니다.",
                survey_id=SURVEY, raw_code="T", survey_code="T",
                facility_label="상수관로시설", target_layer="WTL_PIPE_PS",
                facility_identity="WTL_PIPE_PS::상수관로시설",
                storage_field="ftr_cde", resolved_storage_code="",
                depth_enabled=False,
            )
        with patch('geoflow_connector.ui.survey_panel.QgsMessageLog.logMessage') as log:
            self.assertIsNone(popup.run(start_selected))
        message = log.call_args.args[0]
        self.assertIn('"error_message":', message)
        self.assertIn('"traceback":', message)
        self.assertIn('"facility_identity":', message)
        self.assertIn('storage_code_missing', message)
        popup.close(); popup.deleteLater()

    def test_apply_same_code_empty_standard_code_is_ui_validation(self):
        popup = SurveyPanel(self.plugin, self.window); self.select_source(popup); popup.render_mapping()
        popup.mappingTable.item(0, 0).setText("  ")
        popup.mappingTable.setCurrentCell(0, 0)
        popup.apply_same_code()
        self.assertEqual(
            popup.statusLabel.text(),
            "선택한 행에는 표준 code가 없습니다.",
        )
        popup.close(); popup.deleteLater()

    def test_apply_same_code_with_no_matching_points_is_ui_validation(self):
        popup = SurveyPanel(self.plugin, self.window); self.select_source(popup); popup.render_mapping()
        popup.mappingTable.item(0, 0).setText("NOT_IN_POINTS")
        popup.mappingTable.setCurrentCell(0, 0)
        popup.apply_same_code()
        self.assertEqual(
            popup.statusLabel.text(),
            "같은 코드에 적용할 다른 측량점이 없습니다.",
        )
        popup.close(); popup.deleteLater()

    def test_apply_same_code_keeps_basis_and_has_one_signal_connection(self):
        popup = SurveyPanel(self.plugin, self.window); self.select_source(popup); popup.render_mapping()
        combo = popup.mappingTable.cellWidget(0, 1); combo.setCurrentIndex(1)
        popup.mappingTable.setCurrentCell(0, 0)
        with patch("geoflow_connector.ui.survey_panel.write_mappings") as write:
            popup.applyCodeButton.click()
            self.assertEqual(write.call_count, 1)
            self.assertEqual(popup.mappingTable.currentRow(), 0)
            self.assertEqual(popup.mappingTable.item(0, 0).text(), "M")
            popup.applyCodeButton.click()
            self.assertEqual(write.call_count, 2)
        self.assertTrue(popup.applyCodeButton.isEnabled())
        popup.close(); popup.deleteLater()

    def test_storage_and_depth_settings_persist_and_are_project_scoped(self):
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup); popup.render_mapping()
        storage_combo = popup.storageCodeTable.cellWidget(0, 2)
        storage_combo.setCurrentIndex(1)
        popup.storageCodeTable.cellWidget(0, 3).setChecked(True)
        self.assertIn("저장", popup.storageSaveStatus.text())
        popup.saveStorageCodeButton.click()

        settings, _ = popup._storage_settings()
        storage_key = storage_code_key(self.plugin.active_client.base_url, PROJECT)
        identity = str(popup.storageCodeTable.item(
            0, 0).data(Qt.ItemDataRole.UserRole))
        stored = read_storage_codes(settings, storage_key)
        self.assertEqual(stored[identity]["storage_code"], "SA100")
        depth_key = depth_create_key(self.plugin.active_client.base_url, PROJECT)
        self.assertEqual(read_depth_create_targets(settings, depth_key),
                         {("TEST_LAYER", "SA100")})

        reopened = SurveyPanel(self.plugin, self.window)
        self.select_source(reopened); reopened.render_mapping()
        self.assertEqual(reopened.storageCodeTable.cellWidget(0, 2).currentData()[0],
                         "SA100")
        self.assertTrue(reopened.storageCodeTable.cellWidget(0, 3).isChecked())

        project_b = "20000000-0000-4000-8000-000000000099"
        self.plugin.active_context["project_id"] = project_b
        self.plugin.integration_state = lambda: {
            "ready": True, "can_write": True, "project_id": project_b, "epoch": 2,
        }
        isolated = SurveyPanel(self.plugin, self.window)
        isolated.point_rows = list(popup.point_rows)
        isolated.render_mapping()
        self.assertEqual(isolated.storageCodeTable.cellWidget(0, 2).currentIndex(), 0)
        self.assertEqual(isolated.depthCreateList.count(), 0)
        isolated.close(); isolated.deleteLater()
        reopened.close(); reopened.deleteLater()
        popup.close(); popup.deleteLater()




    def test_source_only_delete_preserves_points_and_features(self):
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup)
        base_get = self.plugin.active_client.get_json.side_effect
        connected = {'delete': {'can_delete': False, 'blocked_reason': 'survey_source_in_use',
                              'point_count': 22,
                              'version': 1,
                              'is_active': True, 'child_version_count': 0,
                              'has_archived_object': True,
                              'original_file_key': 'private/source.csv'}}
        current = connected
        self.plugin.active_client.get_json.side_effect = (
            lambda url: current if f'survey-sources/{SOURCE}/' in url else base_get(url))
        self.plugin.active_client.delete_json_detailed.return_value = {
            'source_deleted': True, 'survey_points_deleted': 0,
            'facilities_deleted': 0,
            'deleted_point_count': 0,
            'retained_point_count': 22, 'retained_target_count': 12,
            'deleted_object': False, 'cleanup_pending': True,
        }
        point_rows_before = list(popup.point_rows)
        feature_count = self.layer.featureCount()
        with patch.object(popup, '_confirm_source_only_delete', return_value=True), \
             patch('geoflow_connector.ui.survey_panel.QgsMessageLog.logMessage') as log:
            popup.delete_source()
        self.plugin.active_client.delete_json_detailed.assert_called_with(
            f'https://example.invalid/gis/projects/{PROJECT}/api/'
            f'survey-sources/{SOURCE}/?mode=source_only'
        )
        self.assertEqual(popup.sourceCombo.findData(SOURCE), -1)
        self.assertGreaterEqual(popup.sourceCombo.findData(popup.RETAINED_POINTS), 0)
        self.assertEqual(popup.point_rows, point_rows_before)
        self.assertEqual(self.layer.featureCount(), feature_count)
        self.assertIn('원본 파일 정리가 지연', popup.statusLabel.text())
        messages = [call.args[0] for call in log.call_args_list]
        self.assertTrue(any('[SURVEY_SOURCE_DELETE] survey_count=22' in value
                            for value in messages))
        self.assertTrue(any('[SURVEY_SOURCE_DELETE] server_result={' in value
                            for value in messages))
        self.assertTrue(any(f'[SURVEY_SOURCE_DELETE] project_id={PROJECT}' in value
                            for value in messages))
        self.assertTrue(any('[SURVEY_SOURCE_DELETE] mode=source_only' in value
                            for value in messages))
        self.assertTrue(any('[SURVEY_SOURCE_DELETE] source_deleted=True' in value
                            for value in messages))
        self.assertTrue(any('[SURVEY_SOURCE_DELETE] survey_points_deleted=0' in value
                            for value in messages))
        reopened = SurveyPanel(self.plugin, self.window)
        retained_index = reopened.sourceCombo.findData(reopened.RETAINED_POINTS)
        self.assertGreaterEqual(retained_index, 0)
        self.assertEqual(reopened.sourceCombo.findData(SOURCE), -1)
        reopened.sourceCombo.setCurrentIndex(retained_index)
        reopened.points()
        self.assertEqual([row['id'] for row in reopened.point_rows], [SURVEY])
        self.assertFalse(hasattr(reopened, '_links_by_point'))
        reopened.close(); reopened.deleteLater()
        popup.close(); popup.deleteLater()

    def test_source_only_rejects_server_child_deletion_response(self):
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup)
        base_get = self.plugin.active_client.get_json.side_effect
        allowed = {'delete': {
            'can_delete_source_only': True, 'point_count': 1,
        }}
        self.plugin.active_client.get_json.side_effect = (
            lambda url: allowed if f'survey-sources/{SOURCE}/' in url else base_get(url)
        )
        self.plugin.active_client.delete_json_detailed.return_value = {
            'source_deleted': True, 'survey_points_deleted': 1,
            'facilities_deleted': 0,
        }
        with patch.object(popup, '_confirm_source_only_delete', return_value=True), \
             self.assertRaisesRegex(RuntimeError, 'Survey Point'):
            popup.delete_source()
        self.assertGreaterEqual(popup.sourceCombo.findData(SOURCE), 0)
        self.assertNotIn(SOURCE, popup._deleted_source_ids())
        popup.close(); popup.deleteLater()

    def test_explicit_raw_code_aliases_and_unknown_passes_through(self):
        settings = QSettings(
            str(Path(self.settings_temp.name) / "aliases.ini"),
            QSettings.Format.IniFormat,
        )
        for raw in ("h", "H"):
            self.assertEqual(normalize_raw_code(PROJECT, raw, settings=settings), "H")
        for raw in ("do", "DO", "Do", "D", "d"):
            self.assertEqual(normalize_raw_code(PROJECT, raw, settings=settings), "DO")
        for raw in ("j", "J", "je", "JE"):
            self.assertEqual(normalize_raw_code(PROJECT, raw, settings=settings), "J")
        self.assertEqual(normalize_raw_code(PROJECT, "ZZZ", settings=settings), "ZZZ")
        self.assertEqual(normalize_raw_code(PROJECT, "PE200", settings=settings), "PE200")
        write_code_aliases(
            settings, normalization_key("https://example.invalid", PROJECT),
            {"custom": "X1"},
        )
        self.assertEqual(normalize_raw_code(
            PROJECT, "custom", settings=settings,
            origin="https://example.invalid",
        ), "X1")
        self.assertEqual(normalize_raw_code(
            SOURCE, "custom", settings=settings,
            origin="https://example.invalid",
        ), "CUSTOM")
        key = normalization_key("https://example.invalid", PROJECT)
        rules = read_code_rules(settings, key)
        rules["CUSTOM"]["enabled"] = False
        write_code_rules(settings, key, rules)
        self.assertEqual(normalize_raw_code(
            PROJECT, "custom", settings=settings,
            origin="https://example.invalid",
        ), "CUSTOM")

    def test_l2_rule_is_selected_after_uppercase(self):
        settings = QSettings(
            str(Path(self.settings_temp.name) / "l2.ini"),
            QSettings.Format.IniFormat,
        )
        key = normalization_key("https://example.invalid", PROJECT)
        write_code_rule_rows(settings, key, [
            {"business_type": "WATER", "raw_code": "je",
             "survey_code": "j", "enabled": True},
            {"business_type": "SEWER", "raw_code": "je",
             "survey_code": "s", "enabled": True},
        ])
        self.assertEqual(normalize_raw_code(
            PROJECT, "Je", settings=settings,
            origin="https://example.invalid", business_type="WATER",
        ), "J")
        self.assertEqual(normalize_raw_code(
            PROJECT, "Je", settings=settings,
            origin="https://example.invalid", business_type="SEWER",
        ), "S")
        self.assertEqual(normalize_raw_code(
            PROJECT, "Je", settings=settings,
            origin="https://example.invalid", business_type="OTHER",
        ), "JE")

    def test_unmapped_code_is_preserved_without_automatic_facility(self):
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup)
        popup.point_rows[0]["raw_code"] = "ZZZ"
        popup.render_features()
        self.assertEqual(popup.featureTable.item(0, 3).text(), "")
        self.assertEqual(
            popup.featureTable.item(0, 2).text(),
            "ZZZ",
        )
        self.assertEqual(
            popup.featureTable.item(0, 0).checkState(), Qt.CheckState.Unchecked
        )
        popup.close(); popup.deleteLater()

    def test_simplified_facility_columns_and_confirmed_first_sort(self):
        popup = SurveyPanel(self.plugin, self.window)
        popup._links_by_point = {}
        popup._links_known = True
        popup._mapping_options = lambda: [
            ("SA100", "A시설", "TEST_LAYER"),
            ("SA200", "B시설", "TEST_LAYER"),
        ]
        popup._saved_mappings = lambda: {
            "J": {"facility_id": "TEST_LAYER::A시설",
                  "facility_label": "A시설", "layer": "TEST_LAYER"},
            "F": {"facility_id": "TEST_LAYER::B시설",
                  "facility_label": "B시설", "layer": "TEST_LAYER"},
        }
        base = {**self.plugin.active_client.get_json.side_effect(
            "survey-points/")["results"][0], "longitude": 127.0,
                "latitude": 36.0, "x": 235000.0, "y": 413000.0,
                "z": 49.0}
        popup.point_rows = [
            {**base, "id": "50000000-0000-4000-8000-000000000011",
             "survey_no": "261017-021", "code": "J"},
            {**base, "id": "50000000-0000-4000-8000-000000000012",
             "survey_no": "261017-003", "code": "J"},
            {**base, "id": "50000000-0000-4000-8000-000000000013",
             "survey_no": "261017-012", "code": "F"},
            {**base, "id": "50000000-0000-4000-8000-000000000014",
             "survey_no": "261017-031", "code": "Z"},
        ]
        popup.render_features()
        self.assertEqual(popup.POINT_COLUMNS, (
            "선택", "측량번호", "Code", "확정시설물", "X", "Y", "Z", "생성상태",
        ))
        self.assertEqual(
            [popup.featureTable.item(row, 1).text() for row in range(4)],
            ["261017-003", "261017-021", "261017-012", "261017-031"],
        )
        self.assertEqual(
            [popup.featureTable.item(row, 3).text() for row in range(4)],
            ["A시설", "A시설", "B시설", ""],
        )
        popup.pages.setCurrentWidget(popup.mappingPage)
        popup.mappingTabs.setCurrentWidget(popup.featurePage)
        popup.resize(1040, 680); popup.show(); APP.processEvents()
        artifact = (Path(__file__).resolve().parents[1] / "test_artifacts" /
                    "survey" / "facility-simple-qgis422.png")
        artifact.parent.mkdir(parents=True, exist_ok=True)
        self.assertTrue(popup.grab().save(str(artifact)))
        popup.close(); popup.deleteLater()

    def test_h_alias_is_reference_only_and_never_a_facility(self):
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup)
        popup.point_rows[0]["raw_code"] = "h"
        popup.render_features()
        self.assertEqual(popup.featureTable.item(0, 3).text(), "")
        self.assertEqual(
            popup.featureTable.item(0, 2).text(), "H"
        )
        self.assertEqual(
            popup.featureTable.item(0, 0).checkState(), Qt.CheckState.Unchecked
        )
        popup.close(); popup.deleteLater()

    def test_legacy_raw_mapping_migrates_to_standard_code_without_deletion(self):
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup)
        popup.point_rows[0]["raw_code"] = "je"
        settings, legacy_key = popup._settings()
        write_mappings(settings, legacy_key, {"je": {
            "facility_label": "맨홀", "layer": "TEST_LAYER",
        }})
        popup.render_mapping()
        self.assertIn("J", popup._saved_mappings())
        self.assertIn("je", read_mappings(settings, legacy_key))
        popup.close(); popup.deleteLater()

    def test_ambiguous_legacy_alias_mapping_requires_reselection(self):
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup)
        settings, legacy_key = popup._settings()
        write_mappings(settings, legacy_key, {
            "j": {"facility_label": "제수변", "layer": "WTL_VALV_PS"},
            "je": {"facility_label": "맨홀", "layer": "TEST_LAYER"},
        })
        migrated = popup._migrate_saved_mappings()
        self.assertNotIn("J", migrated)
        self.assertEqual(set(read_mappings(settings, legacy_key)), {"j", "je"})
        popup.close(); popup.deleteLater()

    def test_source_only_delete_keeps_non_link_safety_blockers(self):
        popup = SurveyPanel(self.plugin, self.window)
        self.select_source(popup)
        base_get = self.plugin.active_client.get_json.side_effect
        blocked = {'delete': {
            'can_delete': False, 'blocked_reason': 'invalid_object_key',
            'point_count': 1, 'linked_target_count': 1,
        }}
        self.plugin.active_client.get_json.side_effect = (
            lambda url: blocked if f'survey-sources/{SOURCE}/' in url else base_get(url)
        )
        with patch('geoflow_connector.ui.survey_panel.QMessageBox.information') as info:
            popup.delete_source()
        info.assert_called_once()
        self.plugin.active_client.delete_json_detailed.assert_not_called()
        popup.close(); popup.deleteLater()

    def test_individual_feature_opens_dynamic_form_draft(self):
        get = self.plugin.active_client.get_json.side_effect
        popup = SurveyPanel(self.plugin, self.window); self.select_source(popup); popup.render_mapping()
        combo = popup.mappingTable.cellWidget(0, 1); combo.setCurrentIndex(1)
        popup.mappingTable.setCurrentCell(0, 0); popup.apply_same_code(); popup.render_features()
        popup._storage_codes = lambda: {
            "TEST_LAYER::맨홀": {
                "standard_name": "TEST_LAYER", "storage_field": "ftr_cde",
                "storage_code": "SA100", "code_label": "맨홀",
            }
        }
        popup.featureTable.setCurrentCell(0, 1); before = self.layer.featureCount()
        popup.create_draft()
        self.assertEqual(self.layer.featureCount(), before + 1)
        draft = self.layer.getFeature(self.layer.selectedFeatureIds()[0])
        self.assertEqual(draft["FTR_CDE"], "SA100")
        self.plugin.work.on_feature_selected.assert_called_once()
        with self.assertRaisesRegex(ValueError, '동일 시설물'):
            popup.create_draft()
        self.layer.rollBack(); popup.close(); popup.deleteLater()

    def test_mixed_layer_selection_plan_and_spatial_duplicates_excluded(self):
        get = self.plugin.active_client.get_json.side_effect
        popup = SurveyPanel(self.plugin, self.window); self.select_source(popup)
        other_layer = QgsVectorLayer('Point?crs=EPSG:4326&field=id:string&field=FTR_CDE:string', 'other', 'memory')
        other_layer.setCustomProperty('geoflow/standard_name', 'OTHER_LAYER')
        self.plugin._managed_layers = lambda: [self.layer, other_layer]
        other = {**popup.point_rows[0], 'id': PROJECT, 'source_row_id': 'P002', 'raw_code': 'J'}
        popup.point_rows.append(other)
        popup._saved_mappings = lambda: {'M': {'ftr_cde': 'SA100', 'layer': 'TEST_LAYER'},
                                         'J': {'ftr_cde': 'OTHER', 'layer': 'OTHER_LAYER'}}
        popup._mapping_options = lambda: [('SA100', 'a', 'TEST_LAYER'), ('OTHER', 'b', 'OTHER_LAYER')]
        popup._storage_codes = lambda: {
            'TEST_LAYER::a': {'standard_name': 'TEST_LAYER', 'storage_code': 'SA100'},
            'OTHER_LAYER::b': {'standard_name': 'OTHER_LAYER', 'storage_code': 'OTHER'},
        }
        popup.render_features()
        self.assertTrue(all(popup.featureTable.item(i, 0).checkState() == Qt.CheckState.Checked for i in range(2)))
        count = self.layer.featureCount()
        plan = popup.prepare_selected()
        self.assertEqual(set(plan), {('TEST_LAYER', 'SA100'), ('OTHER_LAYER', 'OTHER')})
        self.assertEqual(count, self.layer.featureCount())
        feature, _ = popup._new_target_feature(self.layer, popup.point_rows[0], 'SA100')
        self.layer.dataProvider().addFeature(feature)
        popup.render_features()
        self.assertEqual(popup.featureTable.item(0, 0).checkState(), Qt.CheckState.Unchecked)
        self.assertEqual(popup.featureTable.item(0, 7).text(), '●')
        self.assertEqual(popup.featureTable.item(0, 7).toolTip(), '생성됨')
        popup.close(); popup.deleteLater()


class ImportUiTests(unittest.TestCase):
    def test_import_button_runs_preflight_without_validation_button(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "one-click.csv"
            path.write_text(
                "id,code,Easting,Northing,DATE\nP1,j230,200000,500000,2026-10-17\n",
                encoding="utf-8",
            )
            submit = Mock(return_value={
                "local_count": 1, "created": 1, "updated": 0,
                "readback_count": 1,
            })
            dialog = ImportDialog(submit=submit)
            dialog.load_file(path)
            dialog.crsSelector.setCrs(
                QgsCoordinateReferenceSystem("EPSG:5186")
            )
            self.assertTrue(dialog.validateButton.isHidden())
            self.assertTrue(dialog.advancedGroup.isHidden())
            self.assertTrue(dialog.createSurveyButton.isEnabled())
            dialog.create_survey()
            submitted = submit.call_args.args[0]["points"][0]
            self.assertEqual(submitted["raw_code"], "j230")
            self.assertEqual(submitted["code"], "J230")
            dialog.deleteLater()

    def test_designer_required_widget_object_names_and_error_path(self):
        required_widget_names = (
            "errorLabel",
            "encodingCombo", "delimiterCombo", "sheetCombo",
            "sourceTypeCombo", "importProfileCombo",
            "saveImportProfileButton", "manageImportProfileButton",
            "coordinateModeCombo", "crsSelector", "rowIdCombo",
            "rawCodeCombo", "xCombo", "yCombo", "zCombo",
            "surveyDateColumnCombo", "pdopCombo", "surveyHeightCombo",
            "solutionInfoCombo", "previewTable", "validateButton",
            "createSurveyButton", "closeButton",
        )
        dialog = ImportDialog()
        self.assertTrue(dialog.sourceTypeLabel.isHidden())
        self.assertTrue(dialog.sourceTypeCombo.isHidden())
        for name in required_widget_names:
            self.assertTrue(hasattr(dialog, name), name)
            self.assertEqual(getattr(dialog, name).objectName(), name)
        with tempfile.TemporaryDirectory() as folder:
            valid = Path(folder) / "object-names.csv"
            valid.write_text(
                "id,Easting,Northing\nP1,200000,500000\n",
                encoding="utf-8",
            )
            dialog.load_file(valid)
            self.assertFalse(dialog.sourceTypeLabel.isHidden())
            self.assertFalse(dialog.sourceTypeCombo.isHidden())
            self.assertEqual(dialog.previewTable.rowCount(), 1)
            self.assertEqual(dialog.rowIdCombo.currentData(), "id")
            self.assertIn("1행을 읽었습니다", dialog.errorLabel.text())

            invalid = Path(folder) / "unsupported.txt"
            invalid.write_text("not a survey file", encoding="utf-8")
            dialog.load_file(invalid)
            self.assertIsNone(dialog.path)
            self.assertTrue(dialog.sourceTypeLabel.isHidden())
            self.assertTrue(dialog.sourceTypeCombo.isHidden())
            self.assertIn("CSV, XLS 또는 XLSX", dialog.errorLabel.text())

            malformed = Path(folder) / "malformed.csv"
            malformed.write_text(
                "id,,Northing\nP1,200000,500000\n", encoding="utf-8"
            )
            dialog.load_file(malformed)
            self.assertTrue(dialog.sourceTypeLabel.isHidden())
            self.assertTrue(dialog.sourceTypeCombo.isHidden())
            self.assertEqual(dialog.previewTable.rowCount(), 0)
            self.assertIn("열 이름", dialog.errorLabel.text())
        dialog.deleteLater()

    def test_twelve_rows_checked_payload_worker_date_and_gnss(self):
        window, layer, plugin = fixture()
        worker2 = '20000000-0000-4000-8000-000000000002'
        plugin.active_client.session_user.update(worker_id=PROJECT, worker_name='로그인 작업자')
        plugin._reference_service.catalog['workers'] = [dict(id=worker2, name='다른 작업자', resolved=True)]
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'twelve.csv'
            path.write_text('이름,코드,Easting,Northing,측량일,PDOP,안테나높이,solution_info,메모\n' + '\n'.join(
                f'P{i},{"h" if i == 2 else "M"},{200000+i},{500000+i},2026-05-20,1.2,2.0,FIX,현장{i}' for i in range(12)), encoding='utf-8')
            original = path.read_bytes()
            submit = Mock(return_value=dict(local_count=10, created=10, updated=0, readback_count=10))
            dialog = ImportDialog(plugin=plugin, submit=submit)
            self.assertEqual(dialog.workerCombo.currentData(), PROJECT)
            dialog.workerCombo.setCurrentIndex(dialog.workerCombo.findData(worker2))
            dialog.load_file(path); dialog.crsSelector.setCrs(QgsCoordinateReferenceSystem('EPSG:5186'))
            original_raw_code = dialog.previewTable.item(2, 4).text()
            self.assertTrue(dialog.previewTable.item(2, 5).flags()
                            & Qt.ItemFlag.ItemIsEditable)
            self.assertFalse(dialog.previewTable.item(2, 4).flags()
                             & Qt.ItemFlag.ItemIsEditable)
            self.assertFalse(dialog.previewTable.item(2, 6).flags()
                             & Qt.ItemFlag.ItemIsEditable)
            dialog.previewTable.item(2, 5).setText("j")
            self.assertEqual(dialog.previewTable.item(2, 5).text(), "J")
            dialog.previewTable.item(0, 0).setCheckState(Qt.CheckState.Unchecked)
            dialog.previewTable.item(1, 0).setCheckState(Qt.CheckState.Unchecked)
            payload = dialog.validate()
            self.assertEqual(len(payload['points']), 10)
            self.assertEqual(dialog.local_metadata['worker_id'], worker2)
            self.assertEqual(dialog.local_metadata['rows'][0]['survey_date'], '2026-05-20')
            self.assertEqual(dialog.local_metadata['rows'][0]['pdop'], 1.2)
            point = payload['points'][0]
            self.assertEqual(point['worker_id'], worker2)
            self.assertEqual(point['survey_date'], '2026-05-20')
            self.assertEqual(point['name'], 'P2')
            self.assertEqual(point['raw_code'], 'h')
            self.assertEqual(original_raw_code, 'h')
            self.assertEqual(point['code'], 'J')
            headers = [dialog.previewTable.horizontalHeaderItem(i).text()
                       for i in range(dialog.previewTable.columnCount())]
            self.assertIn('raw_code', headers)
            self.assertIn('Survey Code ✎', headers)
            self.assertEqual(point['solution_info'], 'FIX')
            self.assertEqual(point['pdop'], 1.2)
            self.assertEqual(point['antenna_height'], 2.0)
            self.assertEqual(point['raw_data'], {'메모': '현장2'})
            dialog.create_survey()
            self.assertEqual(Path(submit.call_args.args[1]).read_bytes(), original)
            self.assertEqual(len(submit.call_args.args[0]['points']), 10)
            dialog.deleteLater()
        window.close(); window.deleteLater()

    def test_duplicates_auto_uncheck_and_manual_override(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'duplicate.csv'
            path.write_text(
                'id,Easting,Northing,DATE\n'
                'P1,1,2,2026-10-17\nP2,3,4,2026-10-17\nP3,5,6,2026-10-17',
                encoding='utf-8')
            old = [dict(source_row_id='P1', raw_x=1.0, raw_y=2.0,
                        survey_date='2026-10-17', source_group_id=PROJECT,
                        source_crs='EPSG:5186'),
                   dict(source_row_id='P2', raw_x=3.1, raw_y=4.0,
                        survey_date='2026-10-17', source_group_id=PROJECT,
                        source_crs='EPSG:5186')]
            dialog = ImportDialog(duplicate_lookup=lambda: (old, 'test scope'))
            dialog.load_file(path); dialog.crsSelector.setCrs(QgsCoordinateReferenceSystem('EPSG:5186'))
            dialog.sourceGroupEdit.setText(PROJECT); dialog.supersedesEdit.setText(SOURCE)
            payload = dialog.validate()
            self.assertEqual([r['state'] for r in dialog.review], ['DUPLICATE', 'POSSIBLE_DUPLICATE', 'NEW'])
            self.assertEqual(dialog.previewTable.item(0, 1).text(), '●')
            self.assertIn('중복', dialog.previewTable.item(0, 1).toolTip())
            self.assertTrue(dialog.previewTable.item(0, 2).font().strikeOut())
            self.assertIn('중복 의심', dialog.previewTable.item(1, 1).toolTip())
            self.assertIn('원시 3', dialog.previewCountLabel.text())
            self.assertIn('조회 2', dialog.previewCountLabel.text())
            self.assertEqual(len(payload['points']), 2)
            dialog.select_rows('new'); self.assertEqual(len(dialog.validate()['points']), 1)
            dialog.select_rows('all'); self.assertEqual(len(dialog.validate()['points']), 3)
            dialog.select_rows('none'); self.assertIsNone(dialog.validate())
            dialog.deleteLater()
    def test_csv_auto_mapping_dms_and_state_preserved(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "측량.csv"
            path.write_text(
                "이름,코드,경도,위도,Z(레벨),측량일\n"
                "P001,M,127-07-12.1234E,36°16'21.3167\"N,12.3,2026-10-17\n",
                encoding="cp949")
            dialog = ImportDialog(); dialog.load_file(path)
            self.assertEqual(dialog.encodingCombo.currentData(), "cp949")
            dialog.coordinateModeCombo.setCurrentIndex(1)
            self.assertEqual(dialog.rowIdCombo.currentData(), "이름")
            self.assertEqual(dialog.rawCodeCombo.currentData(), "코드")
            self.assertEqual(dialog.previewTable.rowCount(), 1)
            payload = dialog.validate(); self.assertEqual(payload["points"][0]["raw_code"], "M")
            self.assertEqual(dialog.rowIdCombo.currentData(), "이름")
            dialog.deleteLater()

    def test_csv_excel_option_visibility_and_real_encoding_choices(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            utf8 = Path(folder) / "utf8.csv"
            utf8.write_text("id,Easting,Northing\nP1,1,2\n", encoding="utf-8")
            cp949 = Path(folder) / "cp949.csv"
            cp949.write_text("이름,Easting,Northing\nP1,1,2\n", encoding="cp949")
            xlsx = Path(folder) / "book.xlsx"
            book = Workbook(); book.active.title = "Data"
            book.active.append(["id", "Easting", "Northing"])
            book.active.append(["P1", 1, 2]); book.save(xlsx); book.close()
            dialog = ImportDialog()
            self.assertNotIn("auto", [
                dialog.encodingCombo.itemData(i)
                for i in range(dialog.encodingCombo.count())
            ])
            dialog.load_file(utf8)
            self.assertEqual(dialog.encodingCombo.currentData(), "utf-8")
            self.assertFalse(dialog.encodingCombo.isHidden())
            self.assertFalse(dialog.delimiterCombo.isHidden())
            self.assertTrue(dialog.sheetCombo.isHidden())
            dialog.load_file(cp949)
            self.assertEqual(dialog.encodingCombo.currentData(), "cp949")
            dialog.load_file(xlsx)
            self.assertTrue(dialog.encodingCombo.isHidden())
            self.assertTrue(dialog.delimiterCombo.isHidden())
            self.assertFalse(dialog.sheetCombo.isHidden())
            self.assertEqual(dialog.sheetCombo.currentText(), "Data")
            dialog.deleteLater()

    def test_import_profile_persists_and_restores_mapping(self):
        with tempfile.TemporaryDirectory() as folder:
            settings_path = str(Path(folder) / "profiles.ini")
            csv_path = Path(folder) / "foif.csv"
            csv_path.write_text(
                "PT,CODE,X,Y,Z,PDOP,ANT_H,SOL,DATE\n"
                "P1,J,408156.62,236398.45,12.3,1.1,2.0,FIX,2026-10-16\n",
                encoding="utf-8",
            )
            factory = lambda: QSettings(settings_path, QSettings.Format.IniFormat)
            with patch("geoflow_connector.ui.survey_import_dialog.QSettings",
                       side_effect=factory):
                first = ImportDialog(); first.load_file(csv_path)
                first.crsSelector.setCrs(QgsCoordinateReferenceSystem("EPSG:5186"))
                selections = {
                    first.rowIdCombo: "PT", first.rawCodeCombo: "CODE",
                    first.xCombo: "Y", first.yCombo: "X", first.zCombo: "Z",
                    first.pdopCombo: "PDOP", first.surveyHeightCombo: "ANT_H",
                    first.solutionInfoCombo: "SOL",
                    first.surveyDateColumnCombo: "DATE",
                }
                for combo, value in selections.items():
                    combo.setCurrentIndex(combo.findData(value))
                with patch(
                    "geoflow_connector.ui.survey_import_dialog.QInputDialog.getText",
                    return_value=("FOIF A90 GNSS", True),
                ):
                    first.save_profile()
                self.assertIn("FOIF A90 GNSS", first._profiles())
                first.deleteLater()

                second = ImportDialog()
                second._reload_profile_combo("FOIF A90 GNSS")
                second._profile_selected()
                second.load_file(csv_path)
                self.assertEqual(second.sourceTypeCombo.currentData(), "GNSS")
                self.assertEqual(second.coordinateModeCombo.currentIndex(), 0)
                self.assertEqual(second.crsSelector.crs().authid(), "EPSG:5186")
                self.assertEqual(second.xCombo.currentData(), "Y")
                self.assertEqual(second.yCombo.currentData(), "X")
                self.assertEqual(second.pdopCombo.currentData(), "PDOP")
                self.assertEqual(second.surveyDateColumnCombo.currentData(), "DATE")
                payload = second.validate()
                self.assertEqual(payload["source_type"], "GNSS")
                self.assertEqual(payload["survey_type"], "GNSS")
                self.assertEqual(payload["points"][0]["survey_date"], "2026-10-16")
                self.assertEqual(payload["points"][0]["pdop"], 1.1)
                with patch(
                    "geoflow_connector.ui.survey_import_dialog.QInputDialog.getItem",
                    side_effect=[("FOIF A90 GNSS", True), ("삭제", True)],
                ), patch(
                    "geoflow_connector.ui.survey_import_dialog.QMessageBox.question",
                    return_value=QMessageBox.StandardButton.Yes,
                ):
                    second.manage_profiles()
                self.assertEqual(second._profiles(), {})
                second.deleteLater()

    def test_file_date_and_preview_only_survey_number_sequences(self):
        with tempfile.TemporaryDirectory() as folder:
            first_path = Path(folder) / "twenty.csv"
            first_path.write_text(
                "id,Easting,Northing,DATE\n" + "\n".join(
                    f"P{i:02d},{200000+i},{500000+i},2026-10-16" for i in range(20)
                ), encoding="utf-8"
            )
            dialog = ImportDialog(); dialog.load_file(first_path)
            dialog.crsSelector.setCrs(QgsCoordinateReferenceSystem("EPSG:5186"))
            payload = dialog.validate()
            self.assertEqual({row["survey_date"] for row in payload["points"]},
                             {"2026-10-16"})
            headers = [dialog.previewTable.horizontalHeaderItem(i).text()
                       for i in range(dialog.previewTable.columnCount())]
            number_column = headers.index("예상 측량번호")
            numbers = [dialog.previewTable.item(i, number_column).text()
                       for i in range(20)]
            self.assertEqual(numbers[0], "261016-001")
            self.assertEqual(numbers[-1], "261016-020")
            self.assertNotIn("survey_no", payload["points"][0])
            self.assertNotIn("type", payload["points"][0])
            self.assertNotIn("x", payload["points"][0])
            self.assertNotIn("y", payload["points"][0])
            self.assertIsNotNone(payload["points"][0]["longitude"])
            self.assertIsNotNone(payload["points"][0]["latitude"])
            dialog.deleteLater()

            next_path = Path(folder) / "next.csv"
            next_path.write_text(
                "id,Easting,Northing,DATE\n" + "\n".join(
                    f"N{i},{210000+i},{510000+i},2026-10-16" for i in range(5)
                ), encoding="utf-8"
            )
            next_dialog = ImportDialog(); next_dialog.existing_points = [
                {"survey_no": f"261016-{i:03d}"} for i in range(1, 21)
            ]
            next_dialog.load_file(next_path)
            next_dialog.crsSelector.setCrs(QgsCoordinateReferenceSystem("EPSG:5186"))
            next_dialog.validate()
            headers = [next_dialog.previewTable.horizontalHeaderItem(i).text()
                       for i in range(next_dialog.previewTable.columnCount())]
            column = headers.index("예상 측량번호")
            self.assertEqual(next_dialog.previewTable.item(0, column).text(), "261016-021")
            self.assertEqual(next_dialog.previewTable.item(4, column).text(), "261016-025")
            self.assertNotIn("survey_no", next_dialog.payload["points"][0])
            next_dialog.deleteLater()

    def test_114_rows_all_selectable_and_payload_is_complete(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "large.csv"
            path.write_text("id,Easting,Northing,DATE\n" + "\n".join(
                f"P{i:03d},{200000+i},{500000+i},2026-10-17"
                for i in range(114)), encoding="utf-8")
            dialog = ImportDialog(); dialog.load_file(path)
            self.assertIsNone(dialog.validate())
            self.assertIn("좌표계", dialog.errorLabel.text())
            dialog.crsSelector.setCrs(QgsCoordinateReferenceSystem("EPSG:5186"))
            payload = dialog.validate()
            self.assertEqual(dialog.previewTable.rowCount(), 114)
            self.assertIn("원시 114", dialog.previewCountLabel.text())
            self.assertEqual(len(payload["points"]), 114)
            dialog.xCombo.setCurrentIndex(0)
            self.assertTrue(dialog.createSurveyButton.isEnabled())
            self.assertIsNone(dialog.payload)
            dialog.deleteLater()

    def test_axis_labels_manual_xy_and_epsg5186_crosscheck(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "equipment.csv"
            path.write_text(
                "이름,코드,X,Y,Z(레벨),측량일\n"
                "P001,M,408156.62,236398.45,12.3,2026-10-17\n",
                encoding="utf-8")
            dialog = ImportDialog(); dialog.load_file(path)
            self.assertEqual(dialog.firstCoordinateLabel.text(), "Easting (필수)")
            self.assertEqual(dialog.secondCoordinateLabel.text(), "Northing (필수)")
            self.assertIsNone(dialog.xCombo.currentData())
            self.assertIsNone(dialog.yCombo.currentData())
            dialog.xCombo.setCurrentIndex(dialog.xCombo.findData("Y"))
            dialog.yCombo.setCurrentIndex(dialog.yCombo.findData("X"))
            dialog.crsSelector.setCrs(QgsCoordinateReferenceSystem("EPSG:5186"))
            payload = dialog.validate()
            self.assertEqual(payload["source_type"], "GNSS")
            self.assertEqual(payload["survey_type"], "GNSS")
            self.assertAlmostEqual(payload["points"][0]["raw_x"], 236398.45)
            self.assertAlmostEqual(payload["points"][0]["raw_y"], 408156.62)
            self.assertNotIn("x", payload["points"][0])
            self.assertNotIn("y", payload["points"][0])
            self.assertNotIn("z", payload["points"][0])
            self.assertNotIn("type", payload["points"][0])
            transform = QgsCoordinateTransform(
                QgsCoordinateReferenceSystem("EPSG:5186"),
                QgsCoordinateReferenceSystem("EPSG:4326"),
                QgsCoordinateTransformContext())
            point = transform.transform(QgsPointXY(236398.45, 408156.62))
            self.assertAlmostEqual(point.x(), 127.4050808, places=6)
            self.assertAlmostEqual(point.y(), 36.2706871, places=6)
            self.assertAlmostEqual(
                payload["points"][0]["longitude"], point.x(), places=9)
            self.assertAlmostEqual(
                payload["points"][0]["latitude"], point.y(), places=9)
            dialog.coordinateModeCombo.setCurrentIndex(1)
            self.assertEqual(dialog.firstCoordinateLabel.text(), "Longitude (필수)")
            self.assertEqual(dialog.secondCoordinateLabel.text(), "Latitude (필수)")
            dialog.deleteLater()

    def test_total_station_source_contract_and_geographic_target_crs(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "total.csv"
            path.write_text(
                "id,longitude,latitude,DATE\n"
                "TS1,127.1,36.2,2026-10-17\n", encoding="utf-8")
            previous_crs = QgsProject.instance().crs()
            try:
                QgsProject.instance().setCrs(
                    QgsCoordinateReferenceSystem("EPSG:5186"))
                dialog = ImportDialog(); dialog.load_file(path)
                dialog.sourceTypeCombo.setCurrentIndex(
                    dialog.sourceTypeCombo.findData("TOTAL"))
                dialog.coordinateModeCombo.setCurrentIndex(1)
                payload = dialog.validate()
                self.assertEqual(payload["source_type"], "TOTAL")
                self.assertEqual(payload["survey_type"], "TOTAL_STATION")
                self.assertEqual(payload["source_crs"], "EPSG:4326")
                self.assertEqual(payload["target_crs"], "EPSG:5186")
                for field in ("survey_no", "type", "x", "y", "z"):
                    self.assertNotIn(field, payload["points"][0])
                dialog.deleteLater()
            finally:
                QgsProject.instance().setCrs(previous_crs)


class IntegrationTests(unittest.TestCase):
    def test_dock_entry_opens_independent_popup(self):
        window, layer, source = fixture()
        class Base:
            def __init__(self):
                self.__dict__.update(source.__dict__)
                self.window = window; self.dock = QDockWidget(window)
                self.stack = QStackedWidget(); self.dock.setWidget(self.stack)
                root = loadUi(str(Path(__file__).resolve().parents[1] / "ui/designer/main_dock.ui"))
                root.moreActionsButton.setMenu(QMenu(root.moreActionsButton))
                self.presentation = SimpleNamespace(root=root)
                self.integration_events = IntegrationEvents(window)
                self.iface.mainWindow = lambda: window
                self.iface.addPluginToMenu = Mock(); self.iface.addToolBarIcon = Mock()
                self.iface.removePluginMenu = Mock(); self.iface.removeToolBarIcon = Mock()
            def initGui(self): pass
            def unload(self): pass
            def _ensure_ui(self): pass
        class Integrated(SurveyIntegrationMixin, Base): pass
        plugin = Integrated(); plugin.initGui(); plugin._ensure_ui(); plugin._survey_action.trigger(); APP.processEvents()
        self.assertTrue(plugin._survey_panel.isVisible())
        self.assertEqual(plugin.stack.indexOf(plugin._survey_panel), -1)
        self.assertFalse(plugin._survey_panel.isModal())
        plugin.unload(); window.close(); window.deleteLater()

    def test_22_row_presign_put_import_and_readback(self):
        window, _layer, plugin = fixture()
        popup = SurveyPanel(plugin, window)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "survey22.csv"
            path.write_text(
                "id,code,Easting,Northing,Z,Antenna Height,DATE\n" + "\n".join(
                    f"P{i:03d},M,{200000+i},{500000+i},{10+i},0,2026-10-16"
                    for i in range(22)), encoding="utf-8")
            dialog = ImportDialog(popup, submit=popup._submit_import)
            dialog.load_file(path); dialog.crsSelector.setCrs(QgsCoordinateReferenceSystem("EPSG:5186"))
            payload = dialog.validate(); source_id = payload["id"]
            server_points = [{"id": f"50000000-0000-4000-8000-{i:012d}",
                              "source_row_id": f"P{i:03d}", "raw_code": "M",
                              "survey_no": f"261016-{i + 1:03d}",
                              "type": "GNSS", "longitude": 127.0,
                              "latitude": 36.0} for i in range(22)]
            def get(url):
                if "survey-sources/" in url:
                    return {"results": [{"id": source_id, "original_file_name": path.name,
                                          "version": 1, "is_active": True}]}
                return {"results": server_points}
            imported = {}
            def post(url, body):
                if body.get("action") == "presign":
                    return {"ok": True, "id": source_id,
                            "presigned_url": "https://storage.example/upload",
                            "object_key": "gis-survey/t/p/survey22.csv",
                            "headers": {"Content-Type": "text/csv",
                                        "x-amz-server-side-encryption": "AES256"}}
                imported.update(body)
                for target, submitted in zip(server_points, body['points']):
                    target.update(submitted)
                    target.update(
                        x=submitted['raw_x'], y=submitted['raw_y'],
                        z=submitted.get('raw_z'),
                    )
                    target.setdefault('name', submitted['source_row_id'])
                    target.setdefault('code', submitted.get('raw_code'))
                return {"source_id": source_id, "source_group_id": PROJECT,
                        "version": 1, "created": 22, "updated": 0,
                        "points": server_points}
            plugin.active_client.get_json.side_effect = get
            plugin.active_client.post_json.side_effect = post
            plugin.active_client.post_json_detailed.side_effect = post
            plugin.active_client.put_presigned_source_file = Mock()
            dialog.create_survey()
            APP.processEvents()
            self.assertEqual(len(imported["points"]), 22)
            self.assertEqual(imported['points'][0]['survey_date'], '2026-10-16')
            self.assertEqual(imported['source_type'], 'GNSS')
            self.assertEqual(imported['survey_type'], 'GNSS')
            self.assertEqual(imported['points'][0]['antenna_height'], 0.0)
            self.assertNotIn('survey_no', imported['points'][0])
            self.assertNotIn('type', imported['points'][0])
            self.assertNotIn('x', imported['points'][0])
            self.assertNotIn('y', imported['points'][0])
            self.assertNotIn('z', imported['points'][0])
            self.assertIsNotNone(imported['points'][0]['longitude'])
            self.assertIsNotNone(imported['points'][0]['latitude'])
            self.assertEqual(popup.point_rows[0]['survey_date'],
                             imported['points'][0]['survey_date'])
            self.assertEqual(popup.point_rows[0]['name'], 'P000')
            self.assertEqual(popup.point_rows[0]['code'], 'M')
            self.assertEqual(popup.point_rows[0]['survey_no'], '261016-001')
            self.assertEqual(popup.point_rows[0]['type'], 'GNSS')
            self.assertEqual(imported["original_file_key"], "gis-survey/t/p/survey22.csv")
            plugin.active_client.put_presigned_source_file.assert_called_once()
            self.assertEqual(plugin.active_client.put_presigned_source_file.call_args.args[3], "PUT")
            self.assertEqual(popup.pointsTable.rowCount(), 22)
            self.assertEqual(popup.pages.currentWidget(), popup.mappingPage)
            self.assertEqual(popup.mappingTabs.currentWidget(), popup.featurePage)
            plugin._pull_clean_snapshot_delta.assert_called_once_with(
                plugin.active_client, "unused", {})
            self.assertIn("확인 22개", dialog.errorLabel.text())
            self.assertFalse(dialog.createSurveyButton.isEnabled())
            dialog.deleteLater()
        popup.close(); popup.deleteLater(); window.close(); window.deleteLater()

    def test_upload_failure_preserves_validated_state(self):
        window, _layer, plugin = fixture()
        popup = SurveyPanel(plugin, window)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "failed.csv"; path.write_text(
                "id,Easting,Northing,DATE\nP1,1,2,2026-10-17\n", encoding="utf-8")
            dialog = ImportDialog(popup, submit=popup._submit_import); dialog.load_file(path)
            dialog.crsSelector.setCrs(QgsCoordinateReferenceSystem("EPSG:5186")); payload = dialog.validate()
            plugin.active_client.post_json.return_value = {"ok": True,
                "presigned_url": "https://storage.example/upload", "object_key": "key",
                "headers": {}}
            plugin.active_client.post_json.side_effect = None
            plugin.active_client.put_presigned_source_file.side_effect = RuntimeError("network")
            signature = dialog._validated_signature; dialog.create_survey()
            self.assertIs(dialog.payload, payload)
            self.assertEqual(dialog._validated_signature, signature)
            self.assertTrue(dialog.createSurveyButton.isEnabled())
            self.assertIn("업로드에 실패", dialog.errorLabel.text())
            dialog.deleteLater()
        popup.close(); popup.deleteLater(); window.close(); window.deleteLater()

    def test_server_rejection_preserves_validated_state(self):
        window, _layer, plugin = fixture(); popup = SurveyPanel(plugin, window)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "rejected.csv"; path.write_text(
                "id,Easting,Northing,DATE\nP1,1,2,2026-10-17\n", encoding="utf-8")
            dialog = ImportDialog(popup, submit=popup._submit_import); dialog.load_file(path)
            dialog.crsSelector.setCrs(QgsCoordinateReferenceSystem("EPSG:5186")); payload = dialog.validate()
            calls = 0
            def post(_url, body):
                nonlocal calls; calls += 1
                if calls == 1:
                    return {"ok": True, "presigned_url": "https://storage.example/upload",
                            "object_key": "key", "headers": {}}
                raise RuntimeError("survey_rejected")
            plugin.active_client.post_json.side_effect = lambda _url, _body: {
                "ok": True, "presigned_url": "https://storage.example/upload",
                "object_key": "key", "headers": {}}
            plugin.active_client.post_json_detailed.side_effect = RuntimeError("survey_rejected")
            plugin.active_client.put_presigned_source_file = Mock()
            signature = dialog._validated_signature; dialog.create_survey()
            self.assertIs(dialog.payload, payload); self.assertEqual(dialog._validated_signature, signature)
            self.assertTrue(dialog.createSurveyButton.isEnabled()); self.assertIn("survey_rejected", dialog.errorLabel.text())
            dialog.deleteLater()
        popup.close(); popup.deleteLater(); window.close(); window.deleteLater()

    def test_xls_and_xlsx_ui_load(self):
        import xlwt
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            xls = Path(folder) / "old.xls"; book = xlwt.Workbook(); page = book.add_sheet("Sheet1")
            for c, v in enumerate(("id", "x", "y")): page.write(0, c, v)
            for c, v in enumerate(("P1", 1, 2)): page.write(1, c, v)
            book.save(str(xls))
            xlsx = Path(folder) / "new.xlsx"; modern = Workbook(); modern.active.append(["id", "x", "y"]); modern.active.append(["P1", 1, 2]); modern.save(xlsx); modern.close()
            dialog = ImportDialog()
            for path in (xls, xlsx):
                dialog.load_file(path); self.assertEqual(dialog.previewTable.rowCount(), 1)
            dialog.deleteLater()


if __name__ == "__main__":
    print("QGIS_RUNTIME=" + Qgis.QGIS_VERSION)
    unittest.main()
