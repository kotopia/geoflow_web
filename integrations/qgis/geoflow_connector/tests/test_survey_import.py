import tempfile
import unittest
from pathlib import Path

from survey.import_file import (read_table, build_import, parse_angle,
                                suggest_columns, detect_csv_encoding)


class SurveyImportTests(unittest.TestCase):
    def test_xls_is_actually_read(self):
        import xlwt
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'legacy.xls'
            book = xlwt.Workbook(); sheet = book.add_sheet('측량')
            for col, value in enumerate(('이름', 'X', 'Y')): sheet.write(0, col, value)
            for col, value in enumerate(('P001', 200000, 500000)): sheet.write(1, col, value)
            book.save(str(path))
            headers, rows = read_table(path, sheet='측량')
            self.assertEqual(headers, ['이름', 'X', 'Y'])
            self.assertEqual(rows[0]['이름'], 'P001')

    def test_encoding_dms_and_column_suggestions(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'korean.csv'
            path.write_text('이름,코드,경도,위도,Z(레벨)\nP1,M,127,36,1\n', encoding='cp949')
            self.assertEqual(detect_csv_encoding(path), 'cp949')
            headers, _rows, metadata = read_table(path, return_metadata=True)
            self.assertEqual(metadata['encoding'], 'cp949')
            self.assertEqual(suggest_columns(headers, geographic=True), {
                'source_row_id': '이름', 'raw_code': '코드', 'longitude': '경도',
                'latitude': '위도', 'raw_z': 'Z(레벨)'})
        self.assertAlmostEqual(parse_angle('36°16\'21.3167"', latitude=True), 36.2725879722)
        self.assertAlmostEqual(parse_angle('127-07-12.1234E'), 127.1200342778)
        self.assertAlmostEqual(parse_angle('S 36 16 21.3167', latitude=True), -36.2725879722)
        self.assertAlmostEqual(parse_angle('36°16′14.4736″', latitude=True), 36.2706871111)
        self.assertAlmostEqual(parse_angle('127°24′18.2909″'), 127.4050808056)

    def test_planar_suggestions_never_assume_bare_xy_axis_meaning(self):
        self.assertEqual(suggest_columns(
            ["이름", "코드", "X", "Y", "Z(레벨)"], geographic=False), {
                "source_row_id": "이름", "raw_code": "코드",
                "raw_x": None, "raw_y": None, "raw_z": "Z(레벨)",
            })
        self.assertEqual(suggest_columns(
            ["이름", "Easting", "Northing"], geographic=False), {
                "source_row_id": "이름", "raw_code": None,
                "raw_x": "Easting", "raw_y": "Northing", "raw_z": None,
            })

    def test_unicode_dms_is_normalized_before_numeric_validation(self):
        rows = [{"id": "P001", "lon": "127°24′18.2909″",
                 "lat": "36°16′14.4736″"}]
        payload = build_import(
            rows, {"source_row_id": "id", "longitude": "lon", "latitude": "lat"},
            original_file_name="dms.csv", source_type="CSV", geographic=True)
        self.assertAlmostEqual(payload["points"][0]["longitude"], 127.4050808056)
        self.assertAlmostEqual(payload["points"][0]["latitude"], 36.2706871111)

    def test_error_keeps_physical_row_number_and_value(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'rows.csv'
            path.write_text('id,x,y\nP1,1,2\n,,\nP2,bad,4\n', encoding='utf-8')
            _, rows = read_table(path)
        with self.assertRaisesRegex(ValueError, "4행 · raw_x · 값 'bad'"):
            build_import(rows, dict(source_row_id='id', raw_x='x', raw_y='y'),
                         original_file_name='rows.csv', source_type='CSV', source_crs='EPSG:5186')

    def test_csv_cp949_delimiters_and_invalid_headers(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'rows.csv'
            for delimiter in (';', '\t'):
                path.write_text(delimiter.join(['번호', 'X', 'Y']) + '\n' + delimiter.join(['001', '1', '2']), encoding='cp949')
                headers, rows = read_table(path, encoding='cp949', delimiter=delimiter)
                self.assertEqual(rows[0]['번호'], '001')
            for content in ('', 'id,,y\nP1,1,2', 'id,id,y\nP1,1,2', 'id,x,y\n'):
                path.write_text(content, encoding='utf-8')
                with self.assertRaises(ValueError):
                    read_table(path)

    def test_xlsx_cached_formula_is_read_as_value(self):
        from openpyxl import Workbook
        from zipfile import ZipFile
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'cached.xlsx'
            book = Workbook(); book.active.append(['id', 'x', 'y'])
            book.active.append(['P1', '=1+1', 3]); book.save(path); book.close()
            with ZipFile(path) as archive:
                members = {name: archive.read(name) for name in archive.namelist()}
            key = 'xl/worksheets/sheet1.xml'
            members[key] = members[key].replace(b'<f>1+1</f><v></v>', b'<f>1+1</f><v>2</v>')
            with ZipFile(path, 'w') as archive:
                for name, data in members.items():
                    archive.writestr(name, data)
            _, rows = read_table(path)
            self.assertEqual(rows[0]['x'], 2)

    def test_csv_raw_coordinates_and_leading_zero_id(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "survey.csv"
            path.write_text("코드,X,Y,Z\n001,200000.12,500000.34,12.3\n", encoding="utf-8-sig")
            headers, rows = read_table(path)
        payload = build_import(rows, dict(source_row_id="코드", raw_x="X", raw_y="Y", raw_z="Z"),
                               original_file_name="survey.csv", source_type="TOTAL", source_crs="EPSG:5186")
        self.assertEqual(headers, ["코드", "X", "Y", "Z"])
        self.assertEqual(payload["points"], [{"source_row_id": "001", "raw_x": 200000.12,
                                            "raw_y": 500000.34, "raw_z": 12.3}])
        self.assertNotIn("longitude", payload["points"][0])
        self.assertNotIn("original_file_key", payload)

    def test_xlsx_and_revision_geographic(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "survey.xlsx"
            book = Workbook(); page = book.active
            page.append(["id", "lon", "lat"]); page.append(["P001", 127.123, 36.456])
            book.save(path); book.close()
            _, rows = read_table(path)
        payload = build_import(rows, dict(source_row_id="id", longitude="lon", latitude="lat"),
            original_file_name="survey.xlsx", source_type="GNSS", geographic=True,
            source_group_id="30000000-0000-4000-8000-000000000001",
            supersedes_id="30000000-0000-4000-8000-000000000002")
        self.assertEqual(payload["points"][0]["longitude"], 127.123)
        self.assertEqual(payload["points"][0]["raw_x"], 127.123)
        self.assertEqual(payload["points"][0]["raw_y"], 36.456)
        self.assertEqual(payload["points"][0]["raw_crs"], "EPSG:4326")
        self.assertIn("supersedes_id", payload)

    def test_xlsx_formula_is_not_silently_read_as_missing_cached_value(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "formula.xlsx"
            book = Workbook(); book.active.append(["id", "x"]); book.active.append(["P1", "=1+1"])
            book.save(path); book.close()
            with self.assertRaises(ValueError):
                read_table(path)

    def test_input_rejections(self):
        args = dict(original_file_name="f.csv", source_type="TOTAL", source_crs="EPSG:5186")
        mapping = dict(source_row_id="id", raw_x="x", raw_y="y")
        for rows in ([{"id": "P1", "x": "nan", "y": 1}],
                     [{"id": "", "x": 1, "y": 2}],
                     [{"id": "P1", "x": 1, "y": 2}] * 2):
            with self.assertRaises(ValueError):
                build_import(rows, mapping, **args)
        with self.assertRaises(ValueError):
            build_import([{"id": "P1", "x": 1, "y": 2}], mapping, **{**args, "source_crs": ""})


if __name__ == "__main__":
    unittest.main()
